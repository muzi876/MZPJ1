"""数据库初始化，以及把三份清单（岗位、文件类型、现有文件）导入 SQLite。"""
import sqlite3
from pathlib import Path

import openpyxl

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = BASE_DIR / "doc_system.db"
# 三份清单位于本程序所在目录的上一级（基本数据）
DATA_DIR = BASE_DIR.parent

SCHEMA = """
CREATE TABLE IF NOT EXISTS settings (
    key TEXT PRIMARY KEY,
    value TEXT
);

CREATE TABLE IF NOT EXISTS positions (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    department TEXT NOT NULL,
    position TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS file_types (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS existing_files (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    doc_no TEXT,
    name TEXT,
    doc_type TEXT,
    orig_dept TEXT,
    cur_dept TEXT,
    filter_result TEXT,
    main_position TEXT,
    relevance TEXT,
    guest_trip TEXT,
    compliance TEXT,
    reason TEXT,
    person_in_charge TEXT
);

CREATE TABLE IF NOT EXISTS documents (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    doc_type TEXT NOT NULL,
    doc_no TEXT,
    title TEXT,
    purpose TEXT,
    scope TEXT,
    regulation TEXT,
    duty TEXT,
    work_requirement TEXT,
    related_records TEXT,
    related_files TEXT,
    department TEXT,
    position TEXT,
    headcount TEXT,
    direct_supervisor TEXT,
    direct_subordinates TEXT,
    indirect_subordinates TEXT,
    qualifications TEXT,
    responsibilities TEXT,
    work_tasks TEXT,
    relationships TEXT,
    innovation TEXT,
    assessment TEXT,
    created_at TEXT DEFAULT (datetime('now','localtime')),
    updated_at TEXT
);

CREATE TABLE IF NOT EXISTS document_items (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    document_id INTEGER NOT NULL,
    section TEXT,
    seq TEXT,
    level INTEGER DEFAULT 1,
    parent_id INTEGER,
    sort_order INTEGER,
    content TEXT,
    item_type TEXT,
    departments TEXT,
    positions TEXT,
    frequency TEXT,
    keywords TEXT,
    created_at TEXT DEFAULT (datetime('now','localtime'))
);

CREATE TABLE IF NOT EXISTS rasic_matrix (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    document_id INTEGER NOT NULL,
    level TEXT NOT NULL,
    activity TEXT NOT NULL,
    role TEXT NOT NULL,
    letter TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS interface_cards (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    document_id INTEGER NOT NULL,
    name TEXT,
    upstream_dept TEXT,
    downstream_dept TEXT,
    input_content TEXT,
    output_content TEXT,
    form TEXT,
    frequency TEXT,
    related_doc_nos TEXT,
    responsible_positions TEXT
);

CREATE TABLE IF NOT EXISTS flowcharts (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    document_id INTEGER NOT NULL,
    diagram_type TEXT,
    source TEXT
);

CREATE TABLE IF NOT EXISTS employee_panoramas (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    person_name TEXT NOT NULL,
    view_type TEXT NOT NULL,
    narrative TEXT,
    cards TEXT,
    source_doc_nos TEXT,
    created_at TEXT DEFAULT (datetime('now','localtime')),
    updated_at TEXT,
    UNIQUE(person_name, view_type)
);
"""


def get_conn():
    conn = sqlite3.connect(DB_PATH, timeout=30)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA busy_timeout=30000")
    return conn


def init_db():
    conn = get_conn()
    conn.executescript(SCHEMA)
    conn.commit()
    # 为旧库补齐新列（若不存在）
    cols = {r["name"] for r in conn.execute("PRAGMA table_info(document_items)")}
    if "department" not in cols:
        conn.execute("ALTER TABLE document_items ADD COLUMN department TEXT")
    if "main_position" not in cols:
        conn.execute("ALTER TABLE document_items ADD COLUMN main_position TEXT")
    if "content_hash" not in cols:
        conn.execute("ALTER TABLE document_items ADD COLUMN content_hash TEXT")
    if "deleted" not in cols:
        conn.execute("ALTER TABLE document_items ADD COLUMN deleted INTEGER DEFAULT 0")
    conn.commit()
    # 为旧库补齐 documents 新列（工作说明书字段）
    dcols = {r["name"] for r in conn.execute("PRAGMA table_info(documents)")}
    for col in ("department", "position", "headcount", "direct_supervisor",
                "direct_subordinates", "indirect_subordinates", "qualifications",
                "responsibilities", "work_tasks", "relationships", "innovation",
                "assessment"):
        if col not in dcols:
            conn.execute(f"ALTER TABLE documents ADD COLUMN {col} TEXT")
    conn.commit()
    # 为旧库补齐 existing_files 新列（负责人）
    ecols = {r["name"] for r in conn.execute("PRAGMA table_info(existing_files)")}
    if "person_in_charge" not in ecols:
        conn.execute("ALTER TABLE existing_files ADD COLUMN person_in_charge TEXT")
    conn.commit()
    conn.close()


def _cell(v):
    return "" if v is None else str(v).strip()


def import_lists():
    """把三份清单覆盖式导入数据库。"""
    conn = get_conn()

    # 1) 岗位清单
    wb = openpyxl.load_workbook(DATA_DIR / "岗位清单.xlsx", read_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    wb.close()
    conn.execute("DELETE FROM positions")
    for r in rows[1:]:
        if not r or not r[0]:
            continue
        dept, pos = _cell(r[0]), _cell(r[1] if len(r) > 1 else "")
        if dept or pos:
            conn.execute(
                "INSERT INTO positions(department, position) VALUES(?,?)",
                (dept, pos),
            )

    # 2) 文件类型清单
    wb = openpyxl.load_workbook(DATA_DIR / "文件类型清单.xlsx", read_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    wb.close()
    conn.execute("DELETE FROM file_types")
    for r in rows[1:]:
        if r and r[0]:
            conn.execute("INSERT INTO file_types(name) VALUES(?)", (_cell(r[0]),))

    # 3) 现有文件清单
    wb = openpyxl.load_workbook(DATA_DIR / "现有文件清单.xlsx", read_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    wb.close()
    conn.execute("DELETE FROM existing_files")
    for r in rows[1:]:
        if not r or not any(r):
            continue
        padded = list(r) + [None] * 11
        vals = [_cell(x) for x in padded[:11]]
        conn.execute(
            "INSERT INTO existing_files"
            "(doc_no,name,doc_type,orig_dept,cur_dept,filter_result,"
            "main_position,relevance,guest_trip,compliance,reason)"
            " VALUES(?,?,?,?,?,?,?,?,?,?,?)",
            vals,
        )

    # 4) 文件人员匹配表 → 按文件编号回填 person_in_charge
    pic_path = DATA_DIR / "文件人员匹配表.xlsx"
    if pic_path.exists():
        wb2 = openpyxl.load_workbook(pic_path, read_only=True)
        rows2 = list(wb2.active.iter_rows(values_only=True))
        wb2.close()
        for r in rows2[1:]:
            if not r or not r[0]:
                continue
            doc_no = _cell(r[0])
            person = _cell(r[1]) if len(r) > 1 else ""
            conn.execute(
                "UPDATE existing_files SET person_in_charge=? WHERE doc_no=?",
                (person, doc_no),
            )

    conn.commit()
    conn.close()
    return True


def ensure_startup():
    """启动时初始化，若库为空则自动导入三份清单。"""
    init_db()
    conn = get_conn()
    n = conn.execute("SELECT COUNT(*) FROM positions").fetchone()[0]
    conn.close()
    if n == 0:
        import_lists()